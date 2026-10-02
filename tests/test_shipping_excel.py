import unittest
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from modules.shipping.excel import create_shipping_xlsx


class ShippingExcelTests(unittest.TestCase):
    def test_export_contains_only_code_and_price_columns(self):
        content = create_shipping_xlsx([
            {"short_code": "010460123456789321ABCDEFGHIJKLM", "sale_price": Decimal("8490.50")},
        ])

        worksheet = load_workbook(BytesIO(content))["Коды маркировки"]

        self.assertEqual(worksheet.max_column, 2)
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual([worksheet.cell(1, col).value for col in (1, 2)], ["Код маркировки", "Цена продажи"])
        self.assertEqual(worksheet["A2"].value, "010460123456789321ABCDEFGHIJKLM")
        self.assertEqual(worksheet["A2"].number_format, "@")
        self.assertEqual(worksheet["B2"].value, 8490.5)


if __name__ == "__main__":
    unittest.main()
