import tempfile
import unittest
from datetime import date, timedelta
from unittest.mock import patch

from openpyxl import load_workbook

from modules.schedule.config import date_to_str
from modules.schedule.excel import create_schedule_xlsx


class ScheduleExcelTests(unittest.TestCase):
    def test_employee_without_shifts_is_omitted(self):
        week_start = date(2026, 10, 12)
        dates = [week_start + timedelta(days=offset) for offset in range(7)]
        employees = [
            {"employee_id": "working", "full_name": "Работающий сотрудник"},
            {"employee_id": "off", "full_name": "Сотрудник без смен"},
        ]
        schedule = {
            "working": {date_to_str(day): "" for day in dates},
            "off": {date_to_str(day): "" for day in dates},
        }
        schedule["working"][date_to_str(dates[0])] = "12:00"

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as output:
            with patch(
                "modules.schedule.excel.get_schedule_matrix",
                return_value=(employees, dates, schedule, {}),
            ):
                create_schedule_xlsx(week_start, output.name)

            workbook = load_workbook(output.name)
            worksheet = workbook.active

        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(worksheet.cell(row=3, column=1).value, "Работающий сотрудник")
        self.assertEqual(worksheet.cell(row=3, column=2).value, "12:00")
        self.assertNotIn(
            "Сотрудник без смен",
            [worksheet.cell(row=row, column=1).value for row in range(1, worksheet.max_row + 1)],
        )


if __name__ == "__main__":
    unittest.main()
