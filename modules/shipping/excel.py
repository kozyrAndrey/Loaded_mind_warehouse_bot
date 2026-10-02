from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def create_shipping_xlsx(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Коды маркировки"
    worksheet.append(["Код маркировки", "Цена продажи"])
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        worksheet.append([str(row["short_code"]), row["sale_price"]])
        code_cell = worksheet.cell(row=worksheet.max_row, column=1)
        code_cell.number_format = "@"
        worksheet.cell(row=worksheet.max_row, column=2).number_format = '#,##0.00'

    worksheet.column_dimensions["A"].width = 35
    worksheet.column_dimensions["B"].width = 18
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
