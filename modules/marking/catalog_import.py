import csv
import io
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from modules.marking.storage import normalize_gtin


GTIN_HEADERS = {
    "gtin",
    "гтин",
    "штрихкод",
    "штрихкодтовара",
    "кодgtin",
    "gtinтовара",
}
NAME_HEADERS = {
    "название",
    "наименование",
    "наименованиетовара",
    "названиетовара",
    "товар",
    "productname",
    "name",
    "наименованиечестногознака",
    "названиечестногознака",
}
MAX_ROWS = 20_000
MAX_CONSECUTIVE_EMPTY_ROWS = 100


class CatalogImportError(ValueError):
    pass


@dataclass(frozen=True)
class CatalogImportResult:
    rows: list[dict]
    source_name: str
    header_row: int
    gtin_column: str
    name_column: str


def normalize_header(value):
    return re.sub(r"[^0-9a-zа-я]+", "", str(value or "").strip().casefold().replace("ё", "е"))


def column_label(index):
    result = ""
    number = index + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def gtin_cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return format(value, ".0f")
        return str(value)

    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if text.startswith("'"):
        text = text[1:]
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    if re.fullmatch(r"\d+(?:[.,]\d+)?[eE][+-]?\d+", text):
        try:
            number = Decimal(text.replace(",", "."))
            if number == number.to_integral_value():
                text = format(number, "f").split(".", 1)[0]
        except InvalidOperation:
            pass
    return text


def _find_columns(rows):
    for row_index, row in enumerate(rows[:30]):
        headers = [normalize_header(value) for value in row]
        gtin_indexes = [index for index, value in enumerate(headers) if value in GTIN_HEADERS]
        name_indexes = [index for index, value in enumerate(headers) if value in NAME_HEADERS]
        if gtin_indexes and name_indexes:
            return row_index, gtin_indexes[0], name_indexes[0], True

    width = max((len(row) for row in rows[:100]), default=0)
    candidates = []
    for column_index in range(width):
        values = [row[column_index] for row in rows[:100] if column_index < len(row)]
        nonempty = [value for value in values if str(value or "").strip()]
        if not nonempty:
            continue
        valid = 0
        for value in nonempty:
            try:
                normalize_gtin(gtin_cell_text(value))
                valid += 1
            except ValueError:
                pass
        if valid >= 2 and valid / len(nonempty) >= 0.7:
            candidates.append((valid, column_index))
    if not candidates:
        raise CatalogImportError(
            "Не удалось определить колонку GTIN. Добавьте заголовок «GTIN» или «Штрихкод»."
        )

    gtin_index = max(candidates)[1]
    name_candidates = []
    for column_index in range(width):
        if column_index == gtin_index:
            continue
        values = [str(row[column_index] or "").strip() for row in rows[:100] if column_index < len(row)]
        meaningful = [value for value in values if value and not value.isdigit()]
        if meaningful:
            name_candidates.append((len(meaningful), column_index))
    if not name_candidates:
        raise CatalogImportError(
            "Не удалось определить колонку с названием. Добавьте заголовок «Название» или «Наименование»."
        )
    return 0, gtin_index, max(name_candidates)[1], False


def _build_result(rows, source_name):
    if not rows:
        raise CatalogImportError("Таблица пустая.")
    header_index, gtin_index, name_index, has_header = _find_columns(rows)
    data_start = header_index + 1 if has_header else header_index
    products = []
    by_gtin = {}
    errors = []

    for row_number, row in enumerate(rows[data_start:], start=data_start + 1):
        raw_gtin = row[gtin_index] if gtin_index < len(row) else None
        raw_name = row[name_index] if name_index < len(row) else None
        gtin_text = gtin_cell_text(raw_gtin)
        name = str(raw_name or "").strip()
        if not gtin_text and not name:
            continue
        if len(products) >= MAX_ROWS:
            raise CatalogImportError(f"В таблице больше {MAX_ROWS} товарных строк.")
        try:
            gtin = normalize_gtin(gtin_text)
        except ValueError as error:
            errors.append(f"строка {row_number}: {error}")
            continue
        if not name:
            errors.append(f"строка {row_number}: для GTIN {gtin} не указано название.")
            continue
        existing = by_gtin.get(gtin)
        if existing is not None:
            if existing != name:
                errors.append(
                    f"строка {row_number}: GTIN {gtin} уже указан с другим названием «{existing}»."
                )
            continue
        by_gtin[gtin] = name
        products.append({"gtin": gtin, "honest_sign_name": name, "size": ""})

    if errors:
        suffix = "" if len(errors) <= 20 else f"\n…и ещё {len(errors) - 20} ошибок."
        raise CatalogImportError("Ошибки в таблице:\n" + "\n".join(errors[:20]) + suffix)
    if not products:
        raise CatalogImportError("В таблице не найдено ни одной заполненной пары GTIN + название.")

    return CatalogImportResult(
        rows=products,
        source_name=source_name,
        header_row=header_index + 1,
        gtin_column=column_label(gtin_index),
        name_column=column_label(name_index),
    )


def _worksheet_rows(worksheet):
    rows = []
    consecutive_empty = 0
    found_data = False
    for row in worksheet.iter_rows(values_only=True):
        values = list(row)
        is_empty = not any(str(value or "").strip() for value in values)
        rows.append(values)
        if is_empty:
            consecutive_empty += 1
            if found_data and consecutive_empty >= MAX_CONSECUTIVE_EMPTY_ROWS:
                break
        else:
            found_data = True
            consecutive_empty = 0
        if len(rows) > MAX_ROWS + 100:
            raise CatalogImportError(f"На листе больше {MAX_ROWS} строк. Разделите справочник.")
    return rows


def parse_catalog_file(content, filename):
    suffix = str(filename or "").lower().rsplit(".", 1)[-1]
    if suffix == "xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise CatalogImportError(f"Не удалось открыть Excel-файл: {error}") from error
        sheet_errors = []
        for worksheet in workbook.worksheets:
            try:
                rows = _worksheet_rows(worksheet)
                return _build_result(rows, worksheet.title)
            except CatalogImportError as error:
                sheet_errors.append(f"{worksheet.title}: {error}")
        raise CatalogImportError("Не удалось разобрать листы Excel:\n" + "\n".join(sheet_errors[:5]))

    if suffix == "csv":
        decoded = None
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise CatalogImportError("CSV должен быть в кодировке UTF-8 или Windows-1251.")
        try:
            dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return _build_result([list(row) for row in csv.reader(io.StringIO(decoded), dialect)], "CSV")

    raise CatalogImportError("Поддерживаются только файлы .xlsx и .csv.")
